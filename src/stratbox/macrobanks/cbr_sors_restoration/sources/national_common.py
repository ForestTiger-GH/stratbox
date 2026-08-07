from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from stratbox.macrobanks.cbr_sors_restoration.publication import RoundingPolicy
from stratbox.macrobanks.cbr_sors_restoration.sources.common import (
    canonicalize_observations,
    iso_date,
    sha256_file,
    sheet_semantics,
)
from stratbox.macrobanks.cbr_sors_restoration.schema import normalize_text


def parse_national_book(
    path: str | Path,
    as_of_date: str,
    *,
    classifier_id: str,
    source_series: str,
    source_role: str,
    policy: RoundingPolicy,
) -> pd.DataFrame:
    """Parse a historical national workbook in one streaming pass per sheet.

    ``openpyxl`` read-only random ``ws.cell`` access becomes very expensive on the
    century-wide SORS histories.  Sequential rows keep the adapter fast even after
    SME and SME-IE add two more national class books to every run.
    """

    actual = Path(path)
    digest = sha256_file(actual)
    workbook = load_workbook(actual, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    try:
        for ws in workbook.worksheets:
            iterator = ws.iter_rows(values_only=True)
            title_row = next(iterator, ())
            date_row = next(iterator, ())
            title = str((title_row[0] if title_row else None) or ws.title)
            measure, currency = sheet_semantics(title)
            target_index = next(
                (index for index, value in enumerate(date_row) if iso_date(value) == as_of_date),
                None,
            )
            if target_index is None:
                raise ValueError(f'Date {as_of_date} not found in sheet {ws.title!r}')
            for row_no, values in enumerate(iterator, start=3):
                raw_value = values[0] if values else None
                value = values[target_index] if target_index < len(values) else None
                if raw_value is None or value is None:
                    continue
                rows.append({
                    'source_series': source_series,
                    'source_file_actual': actual.name,
                    'source_file_logical': actual.name,
                    'source_sha256': digest,
                    'source_sheet': ws.title,
                    'source_row': row_no,
                    'source_column': target_index + 1,
                    'as_of_date': as_of_date,
                    'classifier_id': classifier_id,
                    'activity_raw': normalize_text(raw_value),
                    'measure': measure,
                    'currency': currency,
                    'value': float(value),
                    'geography_node_id': 'RF',
                    'geography_name': 'РОССИЙСКАЯ ФЕДЕРАЦИЯ',
                    'geography_raw': 'РОССИЙСКАЯ ФЕДЕРАЦИЯ',
                    'geography_kind': 'country_total',
                })
    finally:
        workbook.close()
    out = pd.DataFrame(rows)
    if out.empty:
        raise ValueError(f'No observations parsed from {actual}')
    return canonicalize_observations(out, policy, source_role=source_role)
