from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from stratbox.macrobanks.cbr_sors_restoration.publication import RoundingPolicy
from stratbox.macrobanks.cbr_sors_restoration.schema import normalize_key, normalize_text
from stratbox.macrobanks.cbr_sors_restoration.sources.common import (
    canonicalize_observations,
    iso_date,
    sha256_file,
    sheet_semantics,
)


def parse_regional_totals_history(
    path: str | Path,
    as_of_date: str,
    policy: RoundingPolicy,
    geography_nodes: pd.DataFrame,
) -> pd.DataFrame:
    actual = Path(path)
    digest = sha256_file(actual)
    by_name = {
        normalize_key(row.geography_name): row
        for row in geography_nodes.itertuples(index=False)
    }
    workbook = load_workbook(actual, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    try:
        for ws in workbook.worksheets:
            iterator = ws.iter_rows(values_only=True)
            title_row = next(iterator, ())
            date_row = next(iterator, ())
            measure, currency = sheet_semantics(
                str((title_row[0] if title_row else None) or ws.title)
            )
            target_index = next(
                (
                    index
                    for index, value in enumerate(date_row)
                    if iso_date(value) == as_of_date
                ),
                None,
            )
            if target_index is None:
                raise ValueError(f'Date {as_of_date} not found in sheet {ws.title!r}')
            for row_no, values in enumerate(iterator, start=3):
                raw_name = normalize_text(values[0] if values else None)
                value = values[target_index] if target_index < len(values) else None
                if not raw_name or value is None:
                    continue
                node = by_name.get(normalize_key(raw_name))
                if node is None:
                    raise ValueError(
                        f'01_05_D geography {raw_name!r} is absent from 01_05_A registry'
                    )
                rows.append(
                    {
                        'source_series': '01_05_D',
                        'source_file_actual': actual.name,
                        'source_file_logical': actual.name,
                        'source_sha256': digest,
                        'source_sheet': ws.title,
                        'source_row': row_no,
                        'source_column': target_index + 1,
                        'as_of_date': as_of_date,
                        'classifier_id': 'total_only',
                        'activity_code': 'total',
                        'activity_name': 'ВСЕГО',
                        'activity_raw': 'ВСЕГО',
                        'geography_node_id': str(node.geography_node_id),
                        'geography_name': str(node.geography_name),
                        'geography_raw': raw_name,
                        'geography_kind': str(node.geography_kind),
                        'measure': measure,
                        'currency': currency,
                        'value': float(value),
                    }
                )
    finally:
        workbook.close()
    out = pd.DataFrame(rows)
    if len(out) != 576:
        raise ValueError(f'Expected 576 01_05_D observations, got {len(out)}')
    return canonicalize_observations(
        out,
        policy,
        source_role='OPTIONAL_VALIDATION_AND_HISTORY',
    )
