from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from stratbox.macrobanks.cbr_sors_restoration.publication import RoundingPolicy
from stratbox.macrobanks.cbr_sors_restoration.schema import (
    FD_NAME_TO_CODE,
    normalize_text,
)
from stratbox.macrobanks.cbr_sors_restoration.sources.common import (
    canonicalize_observations,
    iso_date,
    sha256_file,
)


def parse_federal_district_okved2(
    path: str | Path,
    as_of_date: str,
    policy: RoundingPolicy,
    *,
    source_series: str = '01_03_C',
    source_role: str = 'STRICT_FD_OKVED2',
    portfolio_scope: str = 'CORPORATE_TOTAL',
) -> pd.DataFrame:
    actual = Path(path)
    digest = sha256_file(actual)
    workbook = load_workbook(actual, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    seen_date = False
    try:
        for ws in workbook.worksheets:
            title_text = normalize_text(ws.cell(1, 1).value)
            sheet_date = iso_date(title_text)
            if sheet_date:
                seen_date = True
                if sheet_date != as_of_date:
                    raise ValueError(
                        f'FD workbook sheet {ws.title!r} has date {sheet_date}, '
                        f'expected {as_of_date}'
                    )
            lowered = title_text.lower()
            if 'задолж' not in lowered:
                continue
            measure = 'overdue' if 'просроч' in lowered else 'debt'
            header_row = None
            for row_no in range(1, min(ws.max_row, 8) + 1):
                values = [
                    normalize_text(ws.cell(row_no, col).value)
                    for col in range(2, min(ws.max_column, 20) + 1)
                ]
                matches = sum(
                    bool(re.match(r'^[A-Z](?:\.|\s|$)', value))
                    or value.lower().startswith('проч')
                    for value in values
                )
                if matches >= 10:
                    header_row = row_no
                    break
            if header_row is None:
                raise ValueError(f'Cannot locate FD section header in {ws.title!r}')
            headers: dict[int, str] = {}
            for col in range(2, ws.max_column + 1):
                header = normalize_text(ws.cell(header_row, col).value)
                if not header:
                    continue
                headers[col] = (
                    'PUBLISHED_OTHER'
                    if header.lower().startswith('проч')
                    else header.split('.', 1)[0].strip()
                )
            for row_no in range(header_row + 1, ws.max_row + 1):
                fd_name = normalize_text(ws.cell(row_no, 1).value)
                if 'ФЕДЕРАЛЬНЫЙ ОКРУГ' not in fd_name.upper():
                    continue
                fd_code = FD_NAME_TO_CODE.get(fd_name)
                if not fd_code:
                    raise ValueError(f'Unknown federal district name {fd_name!r}')
                for col, section in headers.items():
                    value = ws.cell(row_no, col).value
                    if value is None:
                        continue
                    rows.append(
                        {
                            'source_series': source_series,
                            'source_file_actual': actual.name,
                            'source_file_logical': actual.name,
                            'source_sha256': digest,
                            'source_sheet': ws.title,
                            'source_row': row_no,
                            'source_column': col,
                            'as_of_date': as_of_date,
                            'classifier_id': 'okved2_section',
                            'activity_code': section,
                            'activity_name': normalize_text(
                                ws.cell(header_row, col).value
                            ),
                            'activity_raw': normalize_text(
                                ws.cell(header_row, col).value
                            ),
                            'geography_node_id': fd_code,
                            'geography_name': fd_name,
                            'geography_raw': fd_name,
                            'geography_kind': 'federal_district_total',
                            'measure': measure,
                            'currency': 'total',
                            'value': float(value),
                            'portfolio_scope': portfolio_scope,
                        }
                    )
    finally:
        workbook.close()
    if not seen_date:
        raise ValueError('FD workbook does not contain a recognizable reporting date')
    out = pd.DataFrame(rows)
    if len(out) != 256:
        raise ValueError(f'Expected 256 {source_series} observations, got {len(out)}')
    return canonicalize_observations(
        out,
        policy,
        source_role=source_role,
    )
