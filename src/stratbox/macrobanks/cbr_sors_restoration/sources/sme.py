from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from stratbox.macrobanks.cbr_sors_restoration.publication import RoundingPolicy
from stratbox.macrobanks.cbr_sors_restoration.schema import (
    normalize_class_code,
    normalize_key,
    normalize_text,
)
from stratbox.macrobanks.cbr_sors_restoration.sources.common import (
    canonicalize_observations,
    iso_date,
    sha256_file,
    sheet_semantics,
)
from stratbox.macrobanks.cbr_sors_restoration.sources.national_common import (
    parse_national_book,
)


def _date_column_from_row(values: tuple[object, ...], as_of_date: str) -> int:
    matches = [idx for idx, value in enumerate(values, start=1) if iso_date(value) == as_of_date]
    if not matches:
        raise ValueError(f'Date {as_of_date!r} is absent from workbook header')
    if len(matches) > 1:
        raise ValueError(f'Date {as_of_date!r} appears more than once in workbook header')
    return matches[0]


def parse_sme_national_totals(
    path: str | Path,
    as_of_date: str,
    policy: RoundingPolicy,
) -> pd.DataFrame:
    """Parse 01_11 national SME totals and its SME-IE subset."""

    actual = Path(path)
    digest = sha256_file(actual)
    workbook = load_workbook(actual, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    try:
        for ws in workbook.worksheets:
            iterator = ws.iter_rows(values_only=True)
            title_row = tuple(next(iterator, ()))
            date_row = tuple(next(iterator, ()))
            if not title_row or not date_row:
                raise ValueError(f'01_11 sheet {ws.title!r} is missing header rows')
            measure, currency = sheet_semantics(str(title_row[0] or ws.title))
            column = _date_column_from_row(date_row, as_of_date)
            for row_no, scope in ((3, 'SME'), (4, 'SME_IE')):
                values = tuple(next(iterator, ()))
                label = normalize_text(values[0] if values else None)
                value = values[column - 1] if len(values) >= column else None
                if not label or value is None:
                    raise ValueError(
                        f'01_11 is missing {scope} value in {ws.title!r} at {as_of_date}'
                    )
                rows.append(
                    {
                        'source_series': '01_11',
                        'source_file_actual': actual.name,
                        'source_file_logical': actual.name,
                        'source_sha256': digest,
                        'source_sheet': ws.title,
                        'source_row': row_no,
                        'source_column': column,
                        'as_of_date': as_of_date,
                        'classifier_id': 'total_only',
                        'activity_code': 'total',
                        'activity_name': 'ВСЕГО',
                        'activity_raw': label,
                        'geography_node_id': 'RF',
                        'geography_name': 'РОССИЙСКАЯ ФЕДЕРАЦИЯ',
                        'geography_raw': 'РОССИЙСКАЯ ФЕДЕРАЦИЯ',
                        'geography_kind': 'country_total',
                        'measure': measure,
                        'currency': currency,
                        'value': float(value),
                        'portfolio_scope': scope,
                    }
                )
    finally:
        workbook.close()
    out = pd.DataFrame(rows)
    if len(out) != 12:
        raise ValueError(f'Expected 12 01_11 observations, got {len(out)}')
    return canonicalize_observations(
        out,
        policy,
        source_role='STRICT_NATIONAL_SME_TOTALS',
    )


def parse_sme_national_okved2(
    path: str | Path,
    as_of_date: str,
    policy: RoundingPolicy,
    *,
    portfolio_scope: str,
    source_series: str,
) -> pd.DataFrame:
    if portfolio_scope not in {'SME', 'SME_IE'}:
        raise ValueError(f'Unsupported SME national scope {portfolio_scope!r}')
    out = parse_national_book(
        path,
        as_of_date,
        classifier_id='okved2',
        source_series=source_series,
        source_role=f'STRICT_NATIONAL_OKVED2_{portfolio_scope}',
        policy=policy,
    )
    out['activity_code'] = out['activity_raw'].map(normalize_class_code)
    out = out[
        out['activity_code'].eq('PUBLISHED_OTHER')
        | out['activity_code'].str.fullmatch(r'\d{2}')
    ].copy()
    out['activity_name'] = out['activity_raw']
    out['portfolio_scope'] = portfolio_scope
    if len(out) != 486:
        raise ValueError(f'Expected 486 {source_series} observations, got {len(out)}')
    return out.reset_index(drop=True)


def parse_sme_regional_totals(
    path: str | Path,
    as_of_date: str,
    policy: RoundingPolicy,
    geography_nodes: pd.DataFrame,
    *,
    portfolio_scope: str,
    source_series: str,
) -> pd.DataFrame:
    if portfolio_scope not in {'SME', 'SME_IE'}:
        raise ValueError(f'Unsupported SME regional scope {portfolio_scope!r}')
    actual = Path(path)
    digest = sha256_file(actual)
    by_name = {
        normalize_key(row.geography_name): row
        for row in geography_nodes.itertuples(index=False)
    }
    workbook = load_workbook(actual, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    try:
        for ws in workbook.worksheets:
            iterator = ws.iter_rows(values_only=True)
            title_row = tuple(next(iterator, ()))
            date_row = tuple(next(iterator, ()))
            if not title_row or not date_row:
                raise ValueError(f'{source_series} sheet {ws.title!r} is missing header rows')
            measure, currency = sheet_semantics(str(title_row[0] or ws.title))
            column = _date_column_from_row(date_row, as_of_date)
            for row_no, values in enumerate(iterator, start=3):
                values = tuple(values)
                raw_name = normalize_text(values[0] if values else None)
                value = values[column - 1] if len(values) >= column else None
                if not raw_name or value is None:
                    continue
                node = by_name.get(normalize_key(raw_name))
                if node is None:
                    raise ValueError(
                        f'{source_series} geography {raw_name!r} is absent from 01_05_A registry'
                    )
                rows.append(
                    {
                        'source_series': source_series,
                        'source_file_actual': actual.name,
                        'source_file_logical': actual.name,
                        'source_sha256': digest,
                        'source_sheet': ws.title,
                        'source_row': row_no,
                        'source_column': column,
                        'as_of_date': as_of_date,
                        'classifier_id': 'total_only',
                        'activity_code': 'total',
                        'activity_name': 'ВСЕГО',
                        'activity_raw': 'ВСЕГО',
                        'geography_node_id': str(node.geography_node_id),
                        'geography_name': str(node.geography_name),
                        'geography_raw': raw_name,
                        'geography_kind': str(node.geography_kind),
                        'measure': measure,
                        'currency': currency,
                        'value': float(value),
                        'portfolio_scope': portfolio_scope,
                    }
                )
    finally:
        workbook.close()
    out = pd.DataFrame(rows)
    if len(out) != 576:
        raise ValueError(f'Expected 576 {source_series} observations, got {len(out)}')
    return canonicalize_observations(
        out,
        policy,
        source_role=f'STRICT_REGIONAL_TOTALS_{portfolio_scope}',
    )
