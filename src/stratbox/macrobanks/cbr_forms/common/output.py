"""
Модуль формирует и экспортирует выходные файлы домена ``cbr_forms``.
"""

from __future__ import annotations

from copy import copy
from io import BytesIO
from typing import Mapping

import pandas as pd

from stratbox.base import ioapi as ia
from stratbox.base.ioapi.bytes import write_bytes
from stratbox.base.styles.excel.main import apply_preset
from stratbox.macrobanks.cbr_forms.common.excel_profiles import (
    HIERARCHICAL_STATEMENT_PROFILE,
    STANDARD_PROFILE,
    build_hierarchical_statement_sheets,
)
from stratbox.macrobanks.cbr_forms.common.wide import build_wide_table


EXCEL_META = {
    "creator": "Center for Macroeconomic Analysis and Forecasting (CMAF)",
    "title": "Strategy Box Table Export",
    "category": "CMAF DATA",
}


def export_excel(path: str, df: pd.DataFrame) -> None:
    """
    Функция экспортирует стандартный DataFrame в Excel через общий ``stratbox.ioapi``.
    """
    ia.excel.write_df(
        path,
        df,
        sheet_name="data",
        meta=EXCEL_META,
        style_preset="DEFAULT",
    )
    print(f"[OK] Exported: {path}")


def _set_workbook_meta(wb) -> None:
    """
    Функция записывает стандартные свойства книги Stratbox.
    """
    props = wb.properties
    props.creator = EXCEL_META["creator"]
    props.title = EXCEL_META["title"]
    props.category = EXCEL_META["category"]


def _set_column_widths(ws, widths: Mapping[str, float], *, date_width: float = 15.0) -> None:
    """
    Функция задает предсказуемые ширины колонок специализированной книги.
    """
    for cell in ws[1]:
        header = str(cell.value or "")
        width = widths.get(header)
        if width is None and header and header[0:2].isdigit() and "." in header:
            width = date_width
        if width is not None:
            ws.column_dimensions[cell.column_letter].width = float(width)


def _format_hierarchical_workbook(wb, sheets: dict[str, pd.DataFrame]) -> None:
    """
    Функция применяет оформление иерархической книги без изменения данных.
    """
    freeze = {
        "Обзор": "E2",
        "Полная форма": "F2",
        "Данные": "A2",
    }
    widths = {
        "Раздел": 22,
        "Код": 11,
        "Показатель": 58,
        "Банк": 25,
        "Уровень": 9,
        "Дата": 13,
        "REGN": 10,
        "IndicatorId": 18,
        "Форма": 9,
        "Порядок": 10,
        "РодительскийКод": 18,
        "Measure": 20,
        "Значение": 17,
        "Единица": 17,
    }

    for sheet_name, frame in sheets.items():
        ws = wb[sheet_name]
        apply_preset(ws, "DEFAULT", freeze_panes=freeze[sheet_name])
        if ws.max_row > 1 and ws.max_column > 0:
            ws.auto_filter.ref = ws.dimensions
        _set_column_widths(ws, widths)

        headers = {str(cell.value): cell.column for cell in ws[1]}
        if sheet_name in {"Обзор", "Полная форма"}:
            first_date_col = 5 if sheet_name == "Обзор" else 6
            for row in ws.iter_rows(min_row=2, min_col=first_date_col, max_col=ws.max_column):
                for cell in row:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0"

        if sheet_name == "Данные" and "Значение" in headers:
            value_col = headers["Значение"]
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=value_col)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"

        if sheet_name == "Полная форма":
            indicator_col = headers.get("Показатель")
            level_col = headers.get("Уровень")
            if indicator_col and level_col:
                ws.sheet_properties.outlinePr.summaryBelow = False
                for row_idx in range(2, ws.max_row + 1):
                    level = int(ws.cell(row=row_idx, column=level_col).value or 0)
                    indicator_cell = ws.cell(row=row_idx, column=indicator_col)
                    alignment = copy(indicator_cell.alignment)
                    alignment.indent = min(level, 15)
                    indicator_cell.alignment = alignment
                    ws.row_dimensions[row_idx].outlineLevel = min(level, 7)
                    if level == 0:
                        for cell in ws[row_idx]:
                            font = copy(cell.font)
                            font.bold = True
                            cell.font = font


def export_excel_workbook(path: str, sheets: dict[str, pd.DataFrame]) -> None:
    """
    Функция экспортирует несколько DataFrame в одну оформленную XLSX-книгу.
    """
    from openpyxl import load_workbook

    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

    wb = load_workbook(BytesIO(bio.getvalue()))
    _set_workbook_meta(wb)
    _format_hierarchical_workbook(wb, sheets)

    out = BytesIO()
    wb.save(out)
    write_bytes(path, out.getvalue())
    print(f"[OK] Exported workbook: {path}")


def make_and_export_wide(
    out_path: str,
    df_long: pd.DataFrame,
    df_banks: pd.DataFrame,
    indicator_order: dict[str, int] | None = None,
    date_col: str = "Дата",
    bank_col: str = "Банк",
    indicator_id_col: str = "IndicatorId",
    code_col: str = "Код",
    indicator_col: str = "Показатель",
    value_col: str = "Значение",
) -> pd.DataFrame:
    """
    Функция собирает стандартную wide-таблицу и экспортирует ее в Excel.
    """
    wide_df = build_wide_table(
        df_long=df_long,
        df_banks=df_banks,
        indicator_order=indicator_order,
        date_col=date_col,
        bank_col=bank_col,
        indicator_id_col=indicator_id_col,
        code_col=code_col,
        indicator_col=indicator_col,
        value_col=value_col,
    )
    export_excel(out_path, wide_df)
    return wide_df


def export_form_excel(
    *,
    out_path: str,
    excel_profile: str,
    df_long: pd.DataFrame,
    df_banks: pd.DataFrame,
    model_df: pd.DataFrame,
    indicator_order: dict[str, int] | None = None,
) -> dict[str, pd.DataFrame]:
    """
    Функция выбирает Excel-представление формы по профилю из реестра.

    Стандартный профиль сохраняет прежний одно-листовый формат. Профиль
    ``hierarchical_statement`` формирует обзор, полную форму и canonical long.
    """
    profile = str(excel_profile).strip().lower()
    if profile == STANDARD_PROFILE:
        wide_df = make_and_export_wide(
            out_path=out_path,
            df_long=df_long,
            df_banks=df_banks,
            indicator_order=indicator_order,
        )
        return {"data": wide_df}

    if profile == HIERARCHICAL_STATEMENT_PROFILE:
        sheets = build_hierarchical_statement_sheets(
            df_long=df_long,
            df_banks=df_banks,
            model_df=model_df,
        )
        export_excel_workbook(out_path, sheets)
        return sheets

    raise ValueError(f"Unknown CBR forms Excel profile: {excel_profile!r}")
