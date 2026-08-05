"""Минимальный Excel-экспорт готового pivot-набора ``01_05_A_Debt_corp``."""

from __future__ import annotations

import math
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from stratbox.base.filestore import FileStore
from stratbox.base.runtime import get_filestore
from stratbox.macrobanks.cbr_industries.contracts import (
    Cbr0105ADebtCorpIndustryWorkbookRequest,
    Cbr0105ADebtCorpIndustryWorkbookResult,
    Cbr0105ADebtCorpPivotSetResult,
    Cbr0105ADebtCorpPivotTable,
    Cbr0105ADebtCorpPivotWorkbookRequest,
    Cbr0105ADebtCorpPivotWorkbookResult,
)
from stratbox.macrobanks.cbr_industries.pivots import sanitize_excel_sheet_name
from stratbox.macrobanks.cbr_industries.schema import (
    CBR_0105A_DEBT_CORP_SHEET_SPECS,
    CBR_0105A_DEBT_CORP_UNIT_NAME_RU,
)


_HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
_HEADER_FONT = Font(bold=True)
_THIN_GRAY = Side(style="thin", color="B7C3CF")
_HEADER_BORDER = Border(bottom=_THIN_GRAY)


def _normalize_xlsx_path(path: str) -> str:
    text = str(path).strip()
    if not text:
        raise ValueError("Pivot workbook output path is empty")
    return text if text.lower().endswith(".xlsx") else f"{text}.xlsx"


def _excel_value(value: object) -> object:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, (date, datetime, str, int, float, bool)):
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value
    item = getattr(value, "item", None)
    if callable(item):
        return _excel_value(item())
    return str(value)


def _header_parts(column: object) -> tuple[str, str | None]:
    if isinstance(column, tuple):
        first = "" if column[0] is None else str(column[0])
        second = "" if len(column) < 2 or column[1] is None else str(column[1])
        return first, second
    return str(column), None


def _write_pivot_table_sheet(
    workbook: Workbook,
    *,
    sheet_name: str,
    table,
    freeze_headers: bool,
    enable_auto_filter: bool,
    adjust_column_widths: bool,
) -> None:
    worksheet = workbook.create_sheet(title=sanitize_excel_sheet_name(sheet_name))
    worksheet.sheet_view.showGridLines = False
    dataframe = table.df_table
    columns = list(dataframe.columns)
    descriptor_count = max(0, len(columns) - table.data_columns)
    multi_header = any(isinstance(column, tuple) for column in columns)
    header_rows = 2 if multi_header else 1

    for column_index, column in enumerate(columns, start=1):
        first, second = _header_parts(column)
        worksheet.cell(1, column_index, first)
        if header_rows == 2:
            worksheet.cell(2, column_index, second if second is not None else first)

    data_start_row = header_rows + 1
    for row_offset, row in enumerate(dataframe.itertuples(index=False, name=None), start=0):
        excel_row = data_start_row + row_offset
        for column_index, value in enumerate(row, start=1):
            worksheet.cell(excel_row, column_index, _excel_value(value))

    for row in range(1, header_rows + 1):
        for cell in worksheet[row]:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.border = _HEADER_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if header_rows == 2:
        worksheet.row_dimensions[1].height = 24
        worksheet.row_dimensions[2].height = 24
    else:
        worksheet.row_dimensions[1].height = 24

    last_row = max(header_rows, header_rows + len(dataframe))
    last_column = max(1, len(columns))
    if enable_auto_filter and columns:
        worksheet.auto_filter.ref = (
            f"A{header_rows}:{get_column_letter(last_column)}{last_row}"
        )

    if freeze_headers and columns:
        freeze_column = min(descriptor_count + 1, last_column)
        worksheet.freeze_panes = f"{get_column_letter(freeze_column)}{data_start_row}"

    for column_index in range(descriptor_count + 1, last_column + 1):
        for row in range(data_start_row, last_row + 1):
            worksheet.cell(row, column_index).number_format = "#,##0.##"

    if adjust_column_widths:
        for column_index, column in enumerate(columns, start=1):
            header_first, header_second = _header_parts(column)
            sample_lengths = [len(header_first), len(header_second or "")]
            for row in range(data_start_row, min(last_row, data_start_row + 199) + 1):
                value = worksheet.cell(row, column_index).value
                if value is not None:
                    sample_lengths.append(len(str(value)))
            width = max(sample_lengths, default=8) + 2
            if column_index <= descriptor_count:
                width = min(max(width, 12), 42)
            else:
                width = min(max(width, 11), 18)
            worksheet.column_dimensions[get_column_letter(column_index)].width = width


def _write_metadata_sheet(workbook: Workbook, pivot_set: Cbr0105ADebtCorpPivotSetResult) -> None:
    worksheet = workbook.create_sheet(title="_Параметры", index=0)
    rows: list[tuple[str, object]] = [
        ("row_dimension", pivot_set.row_dimension),
        ("column_dimension", pivot_set.column_dimension),
        ("value_columns", ", ".join(pivot_set.value_columns)),
        ("sheet_dimensions", ", ".join(pivot_set.sheet_dimensions)),
        ("source_rows", pivot_set.source_rows),
        ("table_count", pivot_set.table_count),
        ("is_partial", pivot_set.is_partial),
    ]
    rows.extend((f"fixed:{name}", value) for name, value in pivot_set.fixed_dimensions)
    worksheet.append(["Параметр", "Значение"])
    for name, value in rows:
        worksheet.append([name, _excel_value(value)])
    for cell in worksheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _HEADER_BORDER
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:B{worksheet.max_row}"
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 40


def _workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def save_cbr_0105a_debt_corp_pivot_workbook(
    pivot_set: Cbr0105ADebtCorpPivotSetResult,
    request: Cbr0105ADebtCorpPivotWorkbookRequest,
    *,
    filestore: FileStore | None = None,
) -> Cbr0105ADebtCorpPivotWorkbookResult:
    """Сохраняет все pivot-таблицы отдельными листами одной XLSX-книги."""
    if not pivot_set.tables:
        raise ValueError("Pivot set contains no tables")
    store = filestore or get_filestore()
    output_path = _normalize_xlsx_path(request.out_path)
    if store.exists(output_path) and not request.overwrite:
        raise FileExistsError(f"Pivot workbook already exists: {output_path}")

    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        if request.include_metadata_sheet:
            _write_metadata_sheet(workbook, pivot_set)
        for table in pivot_set.tables:
            _write_pivot_table_sheet(
                workbook,
                sheet_name=table.suggested_sheet_name,
                table=table,
                freeze_headers=request.freeze_headers,
                enable_auto_filter=request.enable_auto_filter,
                adjust_column_widths=request.adjust_column_widths,
            )
        content = _workbook_bytes(workbook)
    finally:
        workbook.close()

    store.write_bytes(output_path, content)
    sheet_names = tuple(
        (["_Параметры"] if request.include_metadata_sheet else [])
        + [table.suggested_sheet_name for table in pivot_set.tables]
    )
    return Cbr0105ADebtCorpPivotWorkbookResult(
        output_path=output_path,
        sheet_names=sheet_names,
        sheet_count=len(sheet_names),
        file_size=len(content),
        sha256=sha256(content).hexdigest(),
    )


_CBR_FONT_NAME = "Times New Roman"
_CBR_TITLE_FONT = Font(name=_CBR_FONT_NAME, size=20, color="333333")
_CBR_SUBTITLE_FONT = Font(name=_CBR_FONT_NAME, size=12, color="333333")
_CBR_HEADER_FONT = Font(name=_CBR_FONT_NAME, size=12, bold=True, color="000000")
_CBR_BODY_FONT = Font(name=_CBR_FONT_NAME, size=12, color="000000")
_CBR_BOLD_FONT = Font(name=_CBR_FONT_NAME, size=12, bold=True, color="000000")
_CBR_COUNTRY_FILL = PatternFill(fill_type="solid", fgColor="CEFFFF")
_CBR_WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF")
_CBR_BLACK_SIDE = Side(style="thin", color="000000")
_CBR_CELL_BORDER = Border(
    left=_CBR_BLACK_SIDE,
    right=_CBR_BLACK_SIDE,
    top=_CBR_BLACK_SIDE,
    bottom=_CBR_BLACK_SIDE,
)
_CBR_SUBTITLE_BORDER = Border(bottom=_CBR_BLACK_SIDE)
_CBR_NUMBER_FORMAT = "#,##0;\\-#,##0;0"


def _format_report_date(value: object) -> str:
    parsed = pd.to_datetime(value, errors="raise")
    return parsed.strftime("%d.%m.%Y")


def _industry_sheet_title(
    *,
    industry_code: str,
    industry_name_ru: str,
    measure_name_ru: str,
    currency_scope: str,
    currency_scope_name_ru: str,
) -> str:
    currency_fragment = ""
    if currency_scope != "total":
        currency_fragment = f" {currency_scope_name_ru.lower()}"
    if industry_code == "apk":
        subject_fragment = f"по отраслевому агрегату «{industry_name_ru}»"
    elif industry_code == "total":
        subject_fragment = (
            "всего по видам экономической деятельности и отдельным "
            "направлениям использования средств"
        )
    else:
        subject_fragment = (
            f"по виду экономической деятельности «{industry_name_ru}»"
        )
    return (
        f"{measure_name_ru} по кредитам, предоставленным юридическим лицам - "
        f"резидентам и индивидуальным предпринимателям, {subject_fragment}"
        f"{currency_fragment}, в региональном разрезе, "
        f"{CBR_0105A_DEBT_CORP_UNIT_NAME_RU}"
    )


def _industry_period_subtitle(dates: tuple[str, ...]) -> str:
    if not dates:
        raise ValueError("Industry workbook contains no report dates")
    if len(dates) == 1:
        return f"Данные по состоянию на {_format_report_date(dates[0])}"
    return (
        "Динамика по состоянию на отчетные даты с "
        f"{_format_report_date(dates[0])} по {_format_report_date(dates[-1])}"
    )


def _table_key(table: Cbr0105ADebtCorpPivotTable) -> tuple[str, str]:
    values = dict(table.sheet_key)
    try:
        return str(values["measure"]), str(values["currency_scope"])
    except KeyError as exc:
        raise ValueError(
            "Industry workbook pivot must be split by measure and currency_scope"
        ) from exc


def _industry_table_dates(table: Cbr0105ADebtCorpPivotTable) -> tuple[str, ...]:
    if table.data_columns < 1:
        raise ValueError("Industry workbook pivot contains no date columns")
    columns = list(table.df_table.columns)
    date_columns = columns[-table.data_columns :]
    if any(isinstance(column, tuple) for column in date_columns):
        raise ValueError("Industry workbook supports exactly one numeric value column")
    return tuple(str(column) for column in date_columns)


def _write_cbr_industry_sheet(
    workbook: Workbook,
    *,
    table: Cbr0105ADebtCorpPivotTable,
    sheet_name: str,
    title: str,
    subtitle: str,
    dates: tuple[str, ...],
) -> None:
    worksheet = workbook.create_sheet(title=sheet_name)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "B4"

    dataframe = table.df_table
    required = {"region_name", "region_kind"}
    missing = sorted(required - set(dataframe.columns))
    if missing:
        raise ValueError(f"Industry workbook table misses region descriptors: {missing}")

    last_column = 1 + len(dates)
    last_column_letter = get_column_letter(last_column)
    worksheet.merge_cells(f"A1:{last_column_letter}1")
    worksheet.merge_cells(f"A2:{last_column_letter}2")

    worksheet["A1"] = title
    if len(dates) < 5:
        title_font = Font(name=_CBR_FONT_NAME, size=12, color="333333")
        title_row_height = 78
    elif len(dates) < 12:
        title_font = Font(name=_CBR_FONT_NAME, size=14, color="333333")
        title_row_height = 47
    else:
        title_font = _CBR_TITLE_FONT
        title_row_height = 31.5
    worksheet["A1"].font = title_font
    worksheet["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )
    worksheet["A1"].fill = _CBR_WHITE_FILL

    worksheet["A2"] = subtitle
    worksheet["A2"].font = _CBR_SUBTITLE_FONT
    worksheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet["A2"].fill = _CBR_WHITE_FILL
    for cell in worksheet[2]:
        cell.border = _CBR_SUBTITLE_BORDER

    worksheet["A3"] = None
    for column_index, report_date in enumerate(dates, start=2):
        worksheet.cell(3, column_index, _format_report_date(report_date))

    for column_index in range(1, last_column + 1):
        cell = worksheet.cell(3, column_index)
        cell.font = _CBR_HEADER_FONT
        cell.fill = _CBR_WHITE_FILL
        cell.border = _CBR_CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    columns = list(dataframe.columns)
    region_name_index = columns.index("region_name")
    region_kind_index = columns.index("region_kind")
    data_start_index = len(columns) - table.data_columns
    for row_offset, row_values in enumerate(
        dataframe.itertuples(index=False, name=None),
        start=4,
    ):
        region_name = _excel_value(row_values[region_name_index])
        region_kind = str(row_values[region_kind_index])
        worksheet.cell(row_offset, 1, region_name)
        for column_index, value in enumerate(
            row_values[data_start_index:],
            start=2,
        ):
            worksheet.cell(row_offset, column_index, _excel_value(value))

        is_country = region_kind == "country_total"
        is_district = region_kind == "federal_district_total"
        row_font = _CBR_BOLD_FONT if is_country or is_district else _CBR_BODY_FONT
        row_fill = _CBR_COUNTRY_FILL if is_country else _CBR_WHITE_FILL
        for column_index in range(1, last_column + 1):
            cell = worksheet.cell(row_offset, column_index)
            cell.font = row_font
            cell.fill = row_fill
            cell.border = _CBR_CELL_BORDER
            cell.alignment = Alignment(
                horizontal="left" if column_index == 1 else "right",
                vertical="center",
                wrap_text=(column_index == 1),
            )
            if column_index > 1:
                cell.number_format = _CBR_NUMBER_FORMAT

        region_text = str(region_name or "")
        worksheet.row_dimensions[row_offset].height = (
            42.2 if is_district or len(region_text) > 58 else 27.75
        )

    worksheet.row_dimensions[1].height = title_row_height
    worksheet.row_dimensions[2].height = 20.85
    worksheet.row_dimensions[3].height = 27.75
    worksheet.column_dimensions["A"].width = 37
    for column_index in range(2, last_column + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 13.5

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:3"
    worksheet.auto_filter.ref = None


def _write_industry_metadata_sheet(
    workbook: Workbook,
    *,
    industry_code: str,
    industry_name_ru: str,
    dates: tuple[str, ...],
) -> None:
    worksheet = workbook.create_sheet(title="_Параметры", index=0)
    rows = [
        ("Серия", "01_05_A_Debt_corp"),
        ("Код отрасли", industry_code),
        ("Наименование отрасли", industry_name_ru),
        ("Первая отчетная дата", dates[0]),
        ("Последняя отчетная дата", dates[-1]),
        ("Количество отчетных дат", len(dates)),
        ("Единица измерения", CBR_0105A_DEBT_CORP_UNIT_NAME_RU),
    ]
    worksheet.append(["Параметр", "Значение"])
    for name, value in rows:
        worksheet.append([name, value])
    for cell in worksheet[1]:
        cell.font = _CBR_HEADER_FONT
        cell.fill = _CBR_COUNTRY_FILL
        cell.border = _CBR_CELL_BORDER
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.font = _CBR_BODY_FONT
            cell.border = _CBR_CELL_BORDER
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 56
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False


def save_cbr_0105a_debt_corp_industry_workbook(
    pivot_set: Cbr0105ADebtCorpPivotSetResult,
    request: Cbr0105ADebtCorpIndustryWorkbookRequest,
    *,
    industry_name_ru: str,
    filestore: FileStore | None = None,
) -> Cbr0105ADebtCorpIndustryWorkbookResult:
    """Сохраняет динамику одной отрасли по регионам в стиле публикации Банка России."""
    if pivot_set.row_dimension != "region_code" or pivot_set.column_dimension != "report_date":
        raise ValueError(
            "Industry workbook requires row_dimension='region_code' and "
            "column_dimension='report_date'"
        )
    if pivot_set.value_columns != ("value",):
        raise ValueError("Industry workbook requires exactly value_columns=('value',)")
    if set(pivot_set.sheet_dimensions) != {"measure", "currency_scope"}:
        raise ValueError(
            "Industry workbook requires sheet_dimensions measure and currency_scope"
        )
    if pivot_set.table_count != 6:
        raise ValueError(
            f"Industry workbook requires six tables, actual={pivot_set.table_count}"
        )

    tables_by_key = {_table_key(table): table for table in pivot_set.tables}
    expected_keys = {
        (spec.measure, spec.currency_scope)
        for spec in CBR_0105A_DEBT_CORP_SHEET_SPECS
    }
    if set(tables_by_key) != expected_keys:
        raise ValueError(
            "Industry workbook table set differs from the six source indicators"
        )

    ordered_specs = tuple(sorted(CBR_0105A_DEBT_CORP_SHEET_SPECS, key=lambda item: item.order))
    first_table = tables_by_key[(ordered_specs[0].measure, ordered_specs[0].currency_scope)]
    dates = _industry_table_dates(first_table)
    if any(_industry_table_dates(table) != dates for table in tables_by_key.values()):
        raise ValueError("Industry workbook sheets contain different report dates")

    store = filestore or get_filestore()
    output_path = _normalize_xlsx_path(request.out_path)
    if store.exists(output_path) and not request.overwrite:
        raise FileExistsError(f"Industry workbook already exists: {output_path}")

    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        if request.include_metadata_sheet:
            _write_industry_metadata_sheet(
                workbook,
                industry_code=request.industry_code,
                industry_name_ru=industry_name_ru,
                dates=dates,
            )
        subtitle = _industry_period_subtitle(dates)
        for spec in ordered_specs:
            table = tables_by_key[(spec.measure, spec.currency_scope)]
            _write_cbr_industry_sheet(
                workbook,
                table=table,
                sheet_name=spec.workbook_sheet_name,
                title=_industry_sheet_title(
                    industry_code=request.industry_code,
                    industry_name_ru=industry_name_ru,
                    measure_name_ru=spec.measure_name_ru,
                    currency_scope=spec.currency_scope,
                    currency_scope_name_ru=spec.currency_scope_name_ru,
                ),
                subtitle=subtitle,
                dates=dates,
            )
        content = _workbook_bytes(workbook)
    finally:
        workbook.close()

    store.write_bytes(output_path, content)
    sheet_names = tuple(
        (["_Параметры"] if request.include_metadata_sheet else [])
        + [spec.workbook_sheet_name for spec in ordered_specs]
    )
    rows_per_sheet = first_table.rows
    return Cbr0105ADebtCorpIndustryWorkbookResult(
        output_path=output_path,
        industry_code=request.industry_code,
        industry_name_ru=industry_name_ru,
        sheet_names=sheet_names,
        sheet_count=len(sheet_names) - (1 if request.include_metadata_sheet else 0),
        dates=dates,
        rows_per_sheet=rows_per_sheet,
        file_size=len(content),
        sha256=sha256(content).hexdigest(),
    )


__all__ = [
    "save_cbr_0105a_debt_corp_industry_workbook",
    "save_cbr_0105a_debt_corp_pivot_workbook",
]
