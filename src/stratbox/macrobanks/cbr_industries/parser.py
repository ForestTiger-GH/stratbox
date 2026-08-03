"""Семантический парсер книг Банка России ``01_05_A_Debt_corp``."""

from __future__ import annotations

import math
import re
import warnings
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from stratbox.macrobanks.cbr_industries.contracts import (
    Cbr0105ADebtCorpDownloadedSource,
    Cbr0105ADebtCorpValidationIssue,
    ParsedCbr0105ADebtCorpFile,
)
from stratbox.macrobanks.cbr_industries.regions import build_cbr_0105a_debt_corp_region_specs
from stratbox.macrobanks.cbr_industries.schema import (
    CBR_0105A_DEBT_CORP_SERIES_CODE,
    CBR_0105A_DEBT_CORP_SHEET_SPECS,
    CBR_0105A_DEBT_CORP_UNIT,
    CBR_0105A_DEBT_CORP_UNIT_NAME_RU,
    clean_cbr_0105a_source_label,
    normalize_cbr_0105a_text,
    resolve_cbr_0105a_debt_corp_industries,
    resolve_cbr_0105a_debt_corp_sheet_spec,
)
from stratbox.macrobanks.cbr_industries.sources import report_date_from_cbr_0105a_filename


_REPORT_DATE_RE = re.compile(r"(?<!\d)(\d{2})[./](\d{2})[./](\d{4})(?!\d)")
_CATEGORY_COLUMNS = (
    "series_code",
    "report_date",
    "sheet_code",
    "measure",
    "measure_name_ru",
    "currency_scope",
    "currency_scope_name_ru",
    "unit",
    "unit_name_ru",
    "region_code",
    "region_name",
    "region_source_name",
    "region_kind",
    "federal_district_name",
    "industry_code",
    "industry_name_ru",
    "industry_source_name",
    "industry_parent_code",
    "source_id",
    "source_url",
    "source_name",
    "source_sha256",
    "source_sheet_name",
    "source_sheet_title",
    "source_column",
)

STREAM_COLUMNS = (
    "series_code",
    "report_date",
    "sheet_code",
    "sheet_order",
    "measure",
    "measure_name_ru",
    "currency_scope",
    "currency_scope_name_ru",
    "unit",
    "unit_name_ru",
    "region_code",
    "region_name",
    "region_source_name",
    "region_kind",
    "federal_district_name",
    "region_order",
    "industry_code",
    "industry_name_ru",
    "industry_source_name",
    "industry_parent_code",
    "industry_hierarchy_level",
    "industry_order",
    "value",
    "source_id",
    "source_url",
    "source_name",
    "source_sha256",
    "source_sheet_name",
    "source_sheet_title",
    "source_row",
    "source_column",
)


def optimize_cbr_0105a_debt_corp_stream_dtypes(df_stream: pd.DataFrame) -> pd.DataFrame:
    """Уплотняет измерения потока для полной месячной истории."""
    stream = df_stream.copy()
    for column in _CATEGORY_COLUMNS:
        if column in stream.columns:
            stream[column] = stream[column].astype("category")
    integer_dtypes = {
        "sheet_order": "int8",
        "region_order": "int16",
        "industry_hierarchy_level": "int8",
        "industry_order": "int8",
        "source_row": "int16",
    }
    for column, dtype in integer_dtypes.items():
        if column in stream.columns:
            stream[column] = stream[column].astype(dtype)
    if "value" in stream.columns:
        stream["value"] = pd.array(stream["value"], dtype="Float64")
    return stream


def _ordered_categories(frames: tuple[pd.DataFrame, ...], column: str) -> list[object]:
    categories: list[object] = []
    seen: set[object] = set()
    for frame in frames:
        if column not in frame.columns:
            continue
        series = frame[column]
        values = (
            series.cat.categories.tolist()
            if isinstance(series.dtype, pd.CategoricalDtype)
            else series.dropna().unique().tolist()
        )
        for value in values:
            if pd.isna(value) or value in seen:
                continue
            seen.add(value)
            categories.append(value)
    return categories


def concat_cbr_0105a_debt_corp_streams(
    frames: Iterable[pd.DataFrame],
) -> pd.DataFrame:
    """Объединяет месячные потоки без распаковки категорий в тяжелый ``object``.

    Pandas сохраняет categorical dtype только при одинаковом наборе категорий во
    всех частях. Поэтому категории сначала выравниваются на небольших справочниках,
    а затем выполняется concat.
    """
    frame_tuple = tuple(frame for frame in frames if frame is not None and not frame.empty)
    if not frame_tuple:
        return pd.DataFrame(columns=STREAM_COLUMNS)

    category_dtypes: dict[str, pd.CategoricalDtype] = {}
    for column in _CATEGORY_COLUMNS:
        if any(column in frame.columns for frame in frame_tuple):
            category_dtypes[column] = pd.CategoricalDtype(
                categories=_ordered_categories(frame_tuple, column),
                ordered=False,
            )

    aligned: list[pd.DataFrame] = []
    for frame in frame_tuple:
        current = frame.copy(deep=False)
        for column, dtype in category_dtypes.items():
            if column in current.columns:
                current[column] = current[column].astype(dtype)
        aligned.append(current)

    stream = pd.concat(aligned, ignore_index=True, copy=False)
    return optimize_cbr_0105a_debt_corp_stream_dtypes(stream)


def _report_date_from_cell(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").replace("\xa0", " ")
    match = _REPORT_DATE_RE.search(text)
    if not match:
        raise ValueError(f"01_05_A_Debt_corp report date is not found in row 2: {value!r}")
    return date(int(match.group(3)), int(match.group(2)), int(match.group(1))).isoformat()


def _coerce_numeric_value(value: object, *, coordinate: str, sheet_name: str) -> int | float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(
            f"Boolean value in 01_05_A_Debt_corp data cell: sheet={sheet_name!r}, cell={coordinate}"
        )
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError(
                f"Non-finite value in 01_05_A_Debt_corp: sheet={sheet_name!r}, cell={coordinate}"
            )
        return int(value) if value.is_integer() else value

    text = str(value).replace("\xa0", " ").strip()
    if not text:
        return None
    normalized = text.replace(" ", "").replace(",", ".")
    try:
        number = Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError(
            "Non-numeric value in 01_05_A_Debt_corp data cell: "
            f"sheet={sheet_name!r}, cell={coordinate}, value={value!r}"
        ) from exc
    if not number.is_finite():
        raise ValueError(
            f"Non-finite value in 01_05_A_Debt_corp: sheet={sheet_name!r}, cell={coordinate}"
        )
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _non_empty_industry_headers(
    sheet_rows: list[tuple[object, ...]],
    *,
    sheet_name: str,
) -> list[tuple[int, object]]:
    if len(sheet_rows) < 3:
        raise ValueError(f"01_05_A_Debt_corp sheet has fewer than 3 rows: {sheet_name!r}")
    headers: list[tuple[int, object]] = []
    for column, value in enumerate(sheet_rows[2][1:], start=2):
        if normalize_cbr_0105a_text(value):
            headers.append((column, value))
    expected_columns = list(range(2, 2 + len(headers)))
    actual_columns = [column for column, _ in headers]
    if actual_columns != expected_columns:
        raise ValueError(
            "01_05_A_Debt_corp industry columns must be contiguous from column B: "
            f"sheet={sheet_name!r}, columns={actual_columns}"
        )
    return headers


def _region_names(
    sheet_rows: list[tuple[object, ...]],
    *,
    sheet_name: str,
) -> tuple[list[int], list[str]]:
    rows: list[int] = []
    names: list[str] = []
    for row, values in enumerate(sheet_rows[3:], start=4):
        value = values[0] if values else None
        if value is None or not str(value).replace("\xa0", " ").strip():
            continue
        rows.append(row)
        names.append(re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip())
    if not rows:
        raise ValueError(f"01_05_A_Debt_corp contains no regions: sheet={sheet_name!r}")
    if rows != list(range(rows[0], rows[0] + len(rows))):
        raise ValueError(
            "01_05_A_Debt_corp geographic rows are not contiguous: "
            f"sheet={sheet_name!r}, rows={rows[:5]}...{rows[-5:]}"
        )
    return rows, names


def _validate_sheet_region_names(
    expected: list[str] | None,
    actual: list[str],
    *,
    sheet_name: str,
) -> list[str]:
    if expected is None:
        return actual
    normalized_expected = [normalize_cbr_0105a_text(value) for value in expected]
    normalized_actual = [normalize_cbr_0105a_text(value) for value in actual]
    if normalized_expected != normalized_actual:
        raise ValueError(
            "Geographic rows differ between 01_05_A_Debt_corp sheets: "
            f"sheet={sheet_name!r}"
        )
    return expected


def _build_validation_issues(
    df_stream: pd.DataFrame,
    *,
    report_date: str,
    source_name: str,
) -> tuple[Cbr0105ADebtCorpValidationIssue, ...]:
    issues: list[Cbr0105ADebtCorpValidationIssue] = []
    missing_count = int(df_stream["value"].isna().sum())
    if missing_count:
        issues.append(
            Cbr0105ADebtCorpValidationIssue(
                code="missing_values",
                severity="warning",
                message="В опубликованной таблице присутствуют пустые значения.",
                count=missing_count,
                details={"report_date": report_date, "source_name": source_name},
            )
        )

    negative_mask = df_stream["value"].notna() & (df_stream["value"] < 0)
    negative_count = int(negative_mask.sum())
    if negative_count:
        issues.append(
            Cbr0105ADebtCorpValidationIssue(
                code="negative_values",
                severity="error",
                message="В остатках задолженности присутствуют отрицательные значения.",
                count=negative_count,
                details={
                    "report_date": report_date,
                    "source_name": source_name,
                    "minimum_value": float(df_stream.loc[negative_mask, "value"].min()),
                },
            )
        )

    pivot = df_stream.pivot_table(
        index=["region_code", "industry_code"],
        columns=["measure", "currency_scope"],
        values="value",
        aggfunc="first",
        dropna=False,
        observed=False,
    )
    for scope in ("rubles", "foreign_currency_and_precious_metals", "total"):
        debt_key = ("debt", scope)
        overdue_key = ("overdue_debt", scope)
        if debt_key not in pivot.columns or overdue_key not in pivot.columns:
            continue
        mask = (
            pivot[overdue_key].notna()
            & pivot[debt_key].notna()
            & (pivot[overdue_key] > pivot[debt_key])
        )
        count = int(mask.sum())
        if count:
            max_excess = float((pivot.loc[mask, overdue_key] - pivot.loc[mask, debt_key]).max())
            issues.append(
                Cbr0105ADebtCorpValidationIssue(
                    code="overdue_exceeds_debt",
                    severity="error",
                    message="Просроченная задолженность превышает общую задолженность.",
                    count=count,
                    details={
                        "report_date": report_date,
                        "source_name": source_name,
                        "currency_scope": scope,
                        "max_excess": max_excess,
                    },
                )
            )

    for measure in ("debt", "overdue_debt"):
        total_key = (measure, "total")
        rubles_key = (measure, "rubles")
        foreign_key = (measure, "foreign_currency_and_precious_metals")
        if not all(key in pivot.columns for key in (total_key, rubles_key, foreign_key)):
            continue
        complete = pivot[[total_key, rubles_key, foreign_key]].notna().all(axis=1)
        gap = (pivot[total_key] - pivot[rubles_key] - pivot[foreign_key]).abs()
        mask = complete & (gap > 1.0)
        count = int(mask.sum())
        if count:
            issues.append(
                Cbr0105ADebtCorpValidationIssue(
                    code="currency_components_do_not_match_total",
                    severity="warning",
                    message=(
                        "Итого расходится с суммой рублевой и валютной частей "
                        "более чем на 1 млн руб."
                    ),
                    count=count,
                    details={
                        "report_date": report_date,
                        "source_name": source_name,
                        "measure": measure,
                        "max_absolute_gap": float(gap.loc[mask].max()),
                    },
                )
            )
    return tuple(issues)


def parse_cbr_0105a_debt_corp_excel_bytes(
    file_bytes: bytes,
    *,
    source_name: str,
    source_id: str | None = None,
    source_url: str | None = None,
    source_sha256: str | None = None,
) -> ParsedCbr0105ADebtCorpFile:
    """Разбирает одну исходную книгу в полный нормализованный поток."""
    filename_date = report_date_from_cbr_0105a_filename(source_name)
    source_identifier = source_id or (
        f"{CBR_0105A_DEBT_CORP_SERIES_CODE}_{filename_date.replace('-', '')}"
    )
    digest = source_sha256 or sha256(file_bytes).hexdigest()

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style.*",
            category=UserWarning,
        )
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)

    try:
        if not workbook.sheetnames:
            raise ValueError("01_05_A_Debt_corp workbook has no worksheets")

        records: list[dict[str, object]] = []
        seen_sheet_codes: set[str] = set()
        parsed_sheet_specs = []
        common_region_names: list[str] | None = None
        common_region_rows: list[int] | None = None
        common_regions = None
        industries = None

        for worksheet in workbook.worksheets:
            sheet_rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
            if len(sheet_rows) < 4:
                raise ValueError(f"01_05_A_Debt_corp sheet is too short: {worksheet.title!r}")
            sheet_title = sheet_rows[0][0] if sheet_rows[0] else None
            sheet_spec = resolve_cbr_0105a_debt_corp_sheet_spec(sheet_title)
            if sheet_spec.code in seen_sheet_codes:
                raise ValueError(
                    "Duplicate semantic sheet in 01_05_A_Debt_corp workbook: "
                    f"code={sheet_spec.code!r}, sheet={worksheet.title!r}"
                )
            seen_sheet_codes.add(sheet_spec.code)
            parsed_sheet_specs.append(sheet_spec)

            sheet_date = _report_date_from_cell(sheet_rows[1][0] if sheet_rows[1] else None)
            if sheet_date != filename_date:
                raise ValueError(
                    "01_05_A_Debt_corp date mismatch between filename and sheet: "
                    f"source_name={source_name!r}, filename_date={filename_date}, "
                    f"sheet={worksheet.title!r}, sheet_date={sheet_date}"
                )

            header_cells = _non_empty_industry_headers(sheet_rows, sheet_name=worksheet.title)
            current_industries = resolve_cbr_0105a_debt_corp_industries(
                [value for _, value in header_cells]
            )
            if industries is None:
                industries = current_industries

            region_rows, region_names = _region_names(sheet_rows, sheet_name=worksheet.title)
            common_region_names = _validate_sheet_region_names(
                common_region_names,
                region_names,
                sheet_name=worksheet.title,
            )
            if common_region_rows is None:
                common_region_rows = region_rows
                common_regions = build_cbr_0105a_debt_corp_region_specs(region_names)
            elif region_rows != common_region_rows:
                raise ValueError(
                    "Geographic row numbers differ between 01_05_A_Debt_corp sheets: "
                    f"sheet={worksheet.title!r}"
                )

            assert common_regions is not None
            for source_row, region in zip(region_rows, common_regions, strict=True):
                for (source_column, source_header), industry in zip(
                    header_cells,
                    current_industries,
                    strict=True,
                ):
                    coordinate = f"{get_column_letter(source_column)}{source_row}"
                    row_values = sheet_rows[source_row - 1]
                    raw_value = (
                        row_values[source_column - 1]
                        if source_column - 1 < len(row_values)
                        else None
                    )
                    value = _coerce_numeric_value(
                        raw_value,
                        coordinate=coordinate,
                        sheet_name=worksheet.title,
                    )
                    records.append(
                        {
                            "series_code": CBR_0105A_DEBT_CORP_SERIES_CODE,
                            "report_date": filename_date,
                            "sheet_code": sheet_spec.code,
                            "sheet_order": sheet_spec.order,
                            "measure": sheet_spec.measure,
                            "measure_name_ru": sheet_spec.measure_name_ru,
                            "currency_scope": sheet_spec.currency_scope,
                            "currency_scope_name_ru": sheet_spec.currency_scope_name_ru,
                            "unit": CBR_0105A_DEBT_CORP_UNIT,
                            "unit_name_ru": CBR_0105A_DEBT_CORP_UNIT_NAME_RU,
                            "region_code": region.code,
                            "region_name": region.canonical_name,
                            "region_source_name": region.source_name,
                            "region_kind": region.region_kind,
                            "federal_district_name": region.federal_district_name,
                            "region_order": region.order,
                            "industry_code": industry.code,
                            "industry_name_ru": industry.canonical_name_ru,
                            "industry_source_name": clean_cbr_0105a_source_label(source_header),
                            "industry_parent_code": industry.parent_code,
                            "industry_hierarchy_level": industry.hierarchy_level,
                            "industry_order": industry.order,
                            "value": value,
                            "source_id": source_identifier,
                            "source_url": source_url,
                            "source_name": source_name,
                            "source_sha256": digest,
                            "source_sheet_name": worksheet.title,
                            "source_sheet_title": str(sheet_title or "").strip(),
                            "source_row": source_row,
                            "source_column": get_column_letter(source_column),
                        }
                    )

        expected_sheet_codes = {spec.code for spec in CBR_0105A_DEBT_CORP_SHEET_SPECS}
        if seen_sheet_codes != expected_sheet_codes:
            raise ValueError(
                "01_05_A_Debt_corp workbook must contain six semantic sheets: "
                f"missing={sorted(expected_sheet_codes - seen_sheet_codes)}, "
                f"unexpected={sorted(seen_sheet_codes - expected_sheet_codes)}"
            )
        assert industries is not None and common_regions is not None

        df_stream = pd.DataFrame.from_records(records, columns=STREAM_COLUMNS)
        df_stream = optimize_cbr_0105a_debt_corp_stream_dtypes(df_stream)
        key_columns = [
            "report_date",
            "measure",
            "currency_scope",
            "region_code",
            "industry_code",
        ]
        duplicate_count = int(df_stream.duplicated(key_columns, keep=False).sum())
        if duplicate_count:
            raise ValueError(
                "Duplicate observations in 01_05_A_Debt_corp stream: "
                f"duplicate_rows={duplicate_count}"
            )

        df_stream = df_stream.sort_values(
            ["report_date", "sheet_order", "region_order", "industry_order"],
            kind="stable",
        ).reset_index(drop=True)
        validation_issues = _build_validation_issues(
            df_stream,
            report_date=filename_date,
            source_name=source_name,
        )

        return ParsedCbr0105ADebtCorpFile(
            source_id=source_identifier,
            source_name=source_name,
            source_url=source_url,
            source_sha256=digest,
            report_date=filename_date,
            sheets=tuple(sorted(parsed_sheet_specs, key=lambda item: item.order)),
            industries=industries,
            regions=common_regions,
            validation_issues=validation_issues,
            df_stream=df_stream,
            rows_stream=int(len(df_stream)),
        )
    finally:
        workbook.close()


def parse_cbr_0105a_debt_corp_source(
    source: Cbr0105ADebtCorpDownloadedSource,
) -> ParsedCbr0105ADebtCorpFile:
    """Разбирает структурированный результат операции скачивания."""
    return parse_cbr_0105a_debt_corp_excel_bytes(
        source.content,
        source_name=source.source_name,
        source_id=source.source_id,
        source_url=source.url,
        source_sha256=source.sha256,
    )


__all__ = [
    "STREAM_COLUMNS",
    "concat_cbr_0105a_debt_corp_streams",
    "optimize_cbr_0105a_debt_corp_stream_dtypes",
    "parse_cbr_0105a_debt_corp_excel_bytes",
    "parse_cbr_0105a_debt_corp_source",
]
