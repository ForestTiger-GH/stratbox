from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from stratbox.macrobanks.cbr_forms.api import _resolve_form_dates
from stratbox.macrobanks.cbr_forms.common.direct_form import build_direct_long
from stratbox.macrobanks.cbr_forms.common.excel_profiles import build_hierarchical_statement_sheets
from stratbox.macrobanks.cbr_forms.common.models import get_model_for, load_models
from stratbox.macrobanks.cbr_forms.common.output import export_form_excel
from stratbox.macrobanks.cbr_forms.forms import form802
from stratbox.macrobanks.cbr_forms.forms.registry import FORM_REGISTRY


def _banks() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"bank": "ВТБ", "regn": 1000, "sort": 1},
            {"bank": "ТЕСТ-БАНК", "regn": 9999, "sort": 2},
        ]
    )


def _raw_snapshot() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"REGN": "1000", "CODE": "1", "total": 100.0, "consolidation_plus": None, "consolidation_minus": None, "intragroup_adjustment": None},
            {"REGN": "1000", "CODE": "2.1.1", "total": 20.0, "consolidation_plus": None, "consolidation_minus": None, "intragroup_adjustment": None},
            {"REGN": "1000", "CODE": "14", "total": 1000.0, "consolidation_plus": None, "consolidation_minus": None, "intragroup_adjustment": None},
            {"REGN": "1000", "CODE": "21", "total": 800.0, "consolidation_plus": None, "consolidation_minus": None, "intragroup_adjustment": None},
            {"REGN": "1000", "CODE": "34", "total": 200.0, "consolidation_plus": None, "consolidation_minus": None, "intragroup_adjustment": None},
        ]
    )


def _long() -> tuple[pd.DataFrame, pd.DataFrame, dict[str, int]]:
    model = get_model_for(load_models(), form="802")
    long, order = build_direct_long(
        date_raw_list=[("01.04.2026", _raw_snapshot())],
        banks_df=_banks(),
        model_df=model,
        spec=form802.DEFAULT_SPEC,
    )
    return long, model, order




def test_form802_uses_quarter_start_reporting_dates_inside_monthly_run_all() -> None:
    dates = _resolve_form_dates(
        entry=FORM_REGISTRY["802"],
        date_from="2026-01-01",
        date_to="2026-08-19",
        default_freq="M",
        default_anchor="start",
    )
    assert [date.strftime("%Y-%m-%d") for date in dates] == [
        "2026-01-01",
        "2026-04-01",
        "2026-07-01",
    ]


def test_standard_form_keeps_global_monthly_schedule() -> None:
    dates = _resolve_form_dates(
        entry=FORM_REGISTRY["101"],
        date_from="2026-01-01",
        date_to="2026-03-01",
        default_freq="M",
        default_anchor="start",
    )
    assert [date.strftime("%Y-%m-%d") for date in dates] == [
        "2026-01-01",
        "2026-02-01",
        "2026-03-01",
    ]


def test_registry_assigns_hierarchical_profile_only_to_802() -> None:
    assert FORM_REGISTRY["802"].excel_profile == "hierarchical_statement"
    assert {
        code: entry.excel_profile
        for code, entry in FORM_REGISTRY.items()
        if code != "802"
    } == {
        "101": "standard",
        "102": "standard",
        "123": "standard",
        "135": "standard",
        "805": "standard",
    }


def test_hierarchical_802_builds_three_expected_views() -> None:
    long, model, _ = _long()
    sheets = build_hierarchical_statement_sheets(
        df_long=long,
        df_banks=_banks(),
        model_df=model,
    )

    assert list(sheets) == ["Обзор", "Полная форма", "Данные"]
    assert len(sheets["Обзор"]) == 34 * 2
    assert len(sheets["Полная форма"]) == 86 * 2
    assert len(sheets["Данные"]) == 86 * 2

    assert sheets["Обзор"].columns[:4].tolist() == ["Раздел", "Код", "Показатель", "Банк"]
    assert sheets["Полная форма"].columns[:5].tolist() == ["Банк", "Раздел", "Уровень", "Код", "Показатель"]

    overview_codes = sheets["Обзор"].loc[
        sheets["Обзор"]["Банк"].eq("ВТБ"),
        "Код",
    ].tolist()
    expected_top_codes = model.loc[model["parent_code"].eq(""), "source_code"].tolist()
    assert overview_codes == expected_top_codes


def test_hierarchical_802_uses_model_tree_for_levels_and_bank_major_full_view() -> None:
    long, model, _ = _long()
    full = build_hierarchical_statement_sheets(
        df_long=long,
        df_banks=_banks(),
        model_df=model,
    )["Полная форма"]

    first_bank = full.iloc[:86]
    second_bank = full.iloc[86:]
    assert first_bank["Банк"].eq("ВТБ").all()
    assert second_bank["Банк"].eq("ТЕСТ-БАНК").all()
    assert first_bank["Код"].tolist() == model["source_code"].tolist()

    levels = first_bank.set_index("Код")["Уровень"]
    assert levels["2"] == 0
    assert levels["2.1"] == 1
    assert levels["2.1.1"] == 2


def test_hierarchical_data_keeps_canonical_columns_and_adds_model_structure() -> None:
    long, model, _ = _long()
    data = build_hierarchical_statement_sheets(
        df_long=long,
        df_banks=_banks(),
        model_df=model,
    )["Данные"]

    for column in long.columns:
        assert column in data.columns
    assert {"Порядок", "РодительскийКод", "Уровень"}.issubset(data.columns)

    row = data[(data["Банк"] == "ВТБ") & (data["Код"] == "2.1.1")].iloc[0]
    assert row["РодительскийКод"] == "2.1"
    assert row["Уровень"] == 2


def test_hierarchical_export_creates_formatted_three_sheet_workbook(tmp_path: Path) -> None:
    long, model, order = _long()
    out_path = tmp_path / "CBR_0409802_LEGACY.xlsx"

    export_form_excel(
        out_path=str(out_path),
        excel_profile="hierarchical_statement",
        df_long=long,
        df_banks=_banks(),
        model_df=model,
        indicator_order=order,
    )

    wb = load_workbook(out_path)
    assert wb.sheetnames == ["Обзор", "Полная форма", "Данные"]
    assert wb["Обзор"].freeze_panes == "E2"
    assert wb["Полная форма"].freeze_panes == "F2"
    assert wb["Данные"].freeze_panes == "A2"
    assert wb["Обзор"].auto_filter.ref
    assert wb["Полная форма"].auto_filter.ref
    assert wb["Данные"].auto_filter.ref

    ws = wb["Полная форма"]
    headers = {cell.value: cell.column for cell in ws[1]}
    code_col = headers["Код"]
    indicator_col = headers["Показатель"]

    row_211 = next(
        row_idx
        for row_idx in range(2, ws.max_row + 1)
        if ws.cell(row=row_idx, column=code_col).value == "2.1.1"
    )
    assert ws.cell(row=row_211, column=indicator_col).alignment.indent == 2
    assert ws.row_dimensions[row_211].outlineLevel == 2
