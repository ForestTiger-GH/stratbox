from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from stratbox.base.filestore import LocalFileStore
from stratbox.macrobanks.cbr_industries.contracts import (
    Cbr0105ADebtCorpPivotRequest,
    Cbr0105ADebtCorpPivotWorkbookRequest,
    Cbr0105ADebtCorpSourceFailure,
    Cbr0105ADebtCorpSourceLink,
    Cbr0105ADebtCorpStreamResult,
    Cbr0105ADebtCorpValidationIssue,
)
from stratbox.macrobanks.cbr_industries.download import (
    try_download_cbr_0105a_debt_corp_source,
)
from stratbox.macrobanks.cbr_industries.operations import (
    build_cbr_0105a_debt_corp_pivot_set,
    save_cbr_0105a_debt_corp_pivot_workbook,
)
from stratbox.macrobanks.cbr_industries.parser import (
    concat_cbr_0105a_debt_corp_streams,
    parse_cbr_0105a_debt_corp_excel_bytes,
)
from stratbox.macrobanks.cbr_industries.pivots import Cbr0105ADebtCorpPivotError
from stratbox.macrobanks.cbr_industries.regions import (
    CBR_0105A_DEBT_CORP_REGION_LAYOUT,
    normalize_cbr_0105a_debt_corp_regions_to_latest,
)
from stratbox.macrobanks.cbr_industries.schema import (
    CBR_0105A_DEBT_CORP_INDUSTRY_SPECS,
)
from stratbox.macrobanks.cbr_industries.sources import (
    discover_cbr_0105a_debt_corp_links_from_html,
)


SHEETS = (
    (
        "rubles",
        "Задолженность по кредитам, предоставленным юридическим лицам - резидентам и индивидуальным предпринимателям в рублях, по видам экономической деятельности и отдельным направлениям использования средств, млн руб.",
        10,
    ),
    (
        " overdue rubles ",
        "Просроченная задолженность по кредитам, предоставленным юридическим лицам - резидентам и индивидуальным предпринимателям в рублях, по видам экономической деятельности и отдельным направлениям использования средств, млн руб.",
        1,
    ),
    (
        "foreign",
        "Задолженность по кредитам, предоставленным юридическим лицам - резидентам и индивидуальным предпринимателям в иностранной валюте и драгоценных металлах, по видам экономической деятельности и отдельным направлениям использования средств, млн руб.",
        5,
    ),
    (
        " overdue foreign ",
        "Просроченная задолженность по кредитам, предоставленным юридическим лицам - резидентам и индивидуальным предпринимателям в иностранной валюте и драгоценных металлах, по видам экономической деятельности и отдельным направлениям использования средств, млн руб.",
        2,
    ),
    (
        "total",
        "Задолженность по кредитам, предоставленным юридическим лицам - резидентам и индивидуальным предпринимателям, по видам экономической деятельности и отдельным направлениям использования средств, млн руб.",
        15,
    ),
    (
        " overdue total ",
        "Просроченная задолженность по кредитам, предоставленным юридическим лицам - резидентам и индивидуальным предпринимателям, по видам экономической деятельности и отдельным направлениям использования средств, млн руб.",
        3,
    ),
)


def _workbook_bytes(
    *,
    report_date: str,
    kemerovo_name: str,
    region_limit: int = 96,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    region_layout = CBR_0105A_DEBT_CORP_REGION_LAYOUT[:region_limit]
    for sheet_name, title, base_value in SHEETS:
        sheet = workbook.create_sheet(sheet_name)
        sheet.cell(1, 1, title)
        sheet.cell(2, 1, f"Задолженность по кредитам по состоянию на {report_date}")
        for column, spec in enumerate(CBR_0105A_DEBT_CORP_INDUSTRY_SPECS, start=2):
            source_header = spec.canonical_name_ru
            if spec.hierarchy_level > 1:
                source_header = "_________\n" * (spec.hierarchy_level - 1) + source_header
            if spec.code in {
                "mining",
                "manufacturing",
                "manufacturing_machinery_equipment",
                "manufacturing_transport_equipment",
                "agriculture_hunting_forestry",
                "construction",
                "transport_communications",
            }:
                source_header += ", из них:"
            sheet.cell(3, column, source_header)

        for row, region in enumerate(region_layout, start=4):
            source_name = (
                kemerovo_name
                if region.code == "0105a_region_081"
                else region.canonical_name
            )
            sheet.cell(row, 1, source_name)
            for column in range(2, 28):
                value = base_value
                if region.code == "0105a_region_081" and column == 2 and base_value == 1:
                    value = "0"
                sheet.cell(row, column, value)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _parsed_pair():
    old = parse_cbr_0105a_debt_corp_excel_bytes(
        _workbook_bytes(
            report_date="01.02.2019",
            kemerovo_name="Кемеровская область",
        ),
        source_name="01_05_A_Debt_corp_20190201.xlsx",
    )
    latest = parse_cbr_0105a_debt_corp_excel_bytes(
        _workbook_bytes(
            report_date="01.06.2026",
            kemerovo_name="Кемеровская область - Кузбасс",
        ),
        source_name="01_05_A_Debt_corp_20260601.xlsx",
    )
    return normalize_cbr_0105a_debt_corp_regions_to_latest([old, latest])


def _stream_result() -> Cbr0105ADebtCorpStreamResult:
    parsed = _parsed_pair()
    stream = concat_cbr_0105a_debt_corp_streams(item.df_stream for item in parsed)
    stream = stream.sort_values(
        ["report_date", "sheet_order", "region_order", "industry_order"],
        kind="stable",
    ).reset_index(drop=True)
    latest = parsed[-1]
    return Cbr0105ADebtCorpStreamResult(
        source_links=(),
        downloaded_sources=(),
        failures=(),
        parsed_files=parsed,
        validation_issues=tuple(
            issue for item in parsed for issue in item.validation_issues
        ),
        df_stream=stream,
        dates=("2019-02-01", "2026-06-01"),
        latest_report_date="2026-06-01",
        latest_discovered_report_date="2026-06-01",
        region_normalization_date="2026-06-01",
        region_normalization_is_latest_discovered=True,
        regions=latest.regions,
        industries=latest.industries,
        rows_stream=len(stream),
    )


def test_discovery_filters_exact_series_and_period() -> None:
    html = b"""
    <a href="/vfs/statistics/BankSector/Loans_to_corporations/01_05_A_Debt_corp_20260401.xlsx">a</a>
    <a href="/vfs/statistics/BankSector/Loans_to_corporations/01_05_A_Debt_corp_20260501.xlsx">b</a>
    <a href="/vfs/statistics/BankSector/Loans_to_corporations/01_05_A_Debt_corp_20260601.xlsx">c</a>
    <a href="/vfs/statistics/BankSector/Loans_to_corporations/01_05_A_Loans_corp_20260601.xlsx">other</a>
    """
    links = discover_cbr_0105a_debt_corp_links_from_html(
        html,
        date_from="2026-05-01",
        date_to="2026-06-01",
    )
    assert [item.report_date for item in links] == ["2026-05-01", "2026-06-01"]
    assert all(item.source_name.startswith("01_05_A_Debt_corp_") for item in links)


def test_parser_recognizes_sheets_and_full_geographic_layout() -> None:
    content = _workbook_bytes(
        report_date="01.06.2026",
        kemerovo_name="Кемеровская область - Кузбасс",
    )
    parsed = parse_cbr_0105a_debt_corp_excel_bytes(
        content,
        source_name="01_05_A_Debt_corp_20260601.xlsx",
    )

    assert parsed.report_date == "2026-06-01"
    assert parsed.rows_stream == 6 * 96 * 26
    assert len(parsed.regions) == 96
    assert [item.code for item in parsed.sheets] == [
        "debt_rubles",
        "overdue_debt_rubles",
        "debt_foreign_currency_and_precious_metals",
        "overdue_debt_foreign_currency_and_precious_metals",
        "debt_total",
        "overdue_debt_total",
    ]
    zero_row = parsed.df_stream[
        (parsed.df_stream["region_source_name"] == "Кемеровская область - Кузбасс")
        & (parsed.df_stream["industry_code"] == "total")
        & (parsed.df_stream["sheet_code"] == "overdue_debt_rubles")
    ].iloc[0]
    assert zero_row["value"] == 0
    assert zero_row["source_column"] == "B"
    assert zero_row["source_row"] == 84


def test_parser_rejects_incomplete_geographic_layout() -> None:
    content = _workbook_bytes(
        report_date="01.06.2026",
        kemerovo_name="Кемеровская область - Кузбасс",
        region_limit=95,
    )
    with pytest.raises(ValueError, match="geographic row count changed"):
        parse_cbr_0105a_debt_corp_excel_bytes(
            content,
            source_name="01_05_A_Debt_corp_20260601.xlsx",
        )


def test_region_names_are_normalized_to_current_local_layout() -> None:
    normalized = _parsed_pair()
    old_stream = normalized[0].df_stream
    row = old_stream[old_stream["region_source_name"] == "Кемеровская область"].iloc[0]
    assert row["region_name"] == "Кемеровская область - Кузбасс"
    old_region = normalized[0].regions[80]
    assert old_region.source_name == "Кемеровская область"
    assert old_region.canonical_name == "Кемеровская область - Кузбасс"


def test_download_reads_valid_source_from_year_cache(tmp_path: Path) -> None:
    content = _workbook_bytes(
        report_date="01.06.2026",
        kemerovo_name="Кемеровская область - Кузбасс",
    )
    source = Cbr0105ADebtCorpSourceLink(
        source_id="01_05_A_Debt_corp_20260601",
        url="https://example.invalid/01_05_A_Debt_corp_20260601.xlsx",
        source_name="01_05_A_Debt_corp_20260601.xlsx",
        report_date="2026-06-01",
    )
    store = LocalFileStore(root=str(tmp_path))
    cache_path = "cache/2026/01_05_A_Debt_corp_20260601.xlsx"
    store.write_bytes(cache_path, content)

    downloaded, failure = try_download_cbr_0105a_debt_corp_source(
        source,
        store=store,
        source_cache_dir="cache",
        refresh=False,
        timeout=1,
        retries=0,
        backoff=0,
        min_bytes_ok=512,
        headers=None,
        plugin_only=False,
    )
    assert failure is None
    assert downloaded is not None
    assert downloaded.from_cache is True
    assert downloaded.cache_path == cache_path
    assert downloaded.content == content


def test_concat_preserves_categorical_dtypes() -> None:
    parsed = _parsed_pair()
    stream = concat_cbr_0105a_debt_corp_streams(item.df_stream for item in parsed)
    assert str(stream["report_date"].dtype) == "category"
    assert str(stream["source_name"].dtype) == "category"
    assert stream["report_date"].astype("object").nunique() == 2


def test_stream_result_ok_accounts_for_validation_errors() -> None:
    result = _stream_result()
    invalid = replace(
        result,
        validation_issues=(
            Cbr0105ADebtCorpValidationIssue(
                code="synthetic_error",
                severity="error",
                message="test",
            ),
        ),
    )
    assert result.ok
    assert not invalid.validation_ok
    assert not invalid.ok


def test_pivot_requires_residual_dimensions_to_be_filtered() -> None:
    result = _stream_result()
    with pytest.raises(Cbr0105ADebtCorpPivotError) as exc_info:
        build_cbr_0105a_debt_corp_pivot_set(
            result.df_stream,
            Cbr0105ADebtCorpPivotRequest(
                row_dimension="region_code",
                column_dimension="report_date",
                sheet_dimensions=("measure", "currency_scope"),
            ),
            stream_result=result,
        )
    assert exc_info.value.code == "PIVOT_RESIDUAL_DIMENSION_VARIATION"
    assert exc_info.value.details["dimension"] == "industry_code"


def test_pivot_builds_six_region_by_date_tables() -> None:
    result = _stream_result()
    filtered = result.df_stream.loc[
        result.df_stream["industry_code"].astype("object") == "total"
    ].copy()
    pivot_set = build_cbr_0105a_debt_corp_pivot_set(
        filtered,
        Cbr0105ADebtCorpPivotRequest(
            row_dimension="region_code",
            column_dimension="report_date",
            sheet_dimensions=("measure", "currency_scope"),
        ),
        stream_result=result,
    )

    assert pivot_set.table_count == 6
    assert [table.suggested_sheet_name for table in pivot_set.tables] == [
        "Задолженность — рубли",
        "Задолженность — инвалюта",
        "Задолженность — итого",
        "Просрочка — рубли",
        "Просрочка — инвалюта",
        "Просрочка — итого",
    ]
    first = pivot_set.tables[0]
    assert first.rows == 96
    assert first.data_columns == 2
    assert list(first.df_table.columns[:4]) == [
        "region_code",
        "region_name",
        "region_kind",
        "federal_district_name",
    ]
    assert list(first.df_table.columns[-2:]) == ["2019-02-01", "2026-06-01"]


def test_pivot_supports_multiple_numeric_value_columns() -> None:
    result = _stream_result()
    filtered = result.df_stream.loc[
        (result.df_stream["region_code"].astype("object") == "0105a_region_001")
        & (result.df_stream["measure"].astype("object") == "debt")
        & (result.df_stream["currency_scope"].astype("object") == "total")
    ].copy()
    filtered["value_copy"] = filtered["value"] * 2
    pivot_set = build_cbr_0105a_debt_corp_pivot_set(
        filtered,
        Cbr0105ADebtCorpPivotRequest(
            row_dimension="industry_code",
            column_dimension="report_date",
            value_columns=("value", "value_copy"),
        ),
        stream_result=result,
    )
    table = pivot_set.tables[0]
    assert table.rows == 26
    assert table.data_columns == 4
    assert any(isinstance(column, tuple) for column in table.df_table.columns)


def test_save_pivot_workbook_writes_all_tables(tmp_path: Path) -> None:
    result = _stream_result()
    filtered = result.df_stream.loc[
        result.df_stream["industry_code"].astype("object") == "total"
    ].copy()
    pivot_set = build_cbr_0105a_debt_corp_pivot_set(
        filtered,
        Cbr0105ADebtCorpPivotRequest(
            row_dimension="region_code",
            column_dimension="report_date",
            sheet_dimensions=("measure", "currency_scope"),
        ),
        stream_result=result,
    )
    store = LocalFileStore(root=str(tmp_path))
    saved = save_cbr_0105a_debt_corp_pivot_workbook(
        pivot_set,
        Cbr0105ADebtCorpPivotWorkbookRequest(
            out_path="result",
            include_metadata_sheet=True,
        ),
        filestore=store,
    )

    assert saved.ok
    assert saved.output_path == "result.xlsx"
    assert saved.sheet_count == 7
    assert saved.file_size > 0
    workbook = load_workbook(tmp_path / "result.xlsx", read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "_Параметры",
            "Задолженность — рубли",
            "Задолженность — инвалюта",
            "Задолженность — итого",
            "Просрочка — рубли",
            "Просрочка — инвалюта",
            "Просрочка — итого",
        ]
        sheet = workbook["Задолженность — рубли"]
        assert sheet["A1"].value == "region_code"
        assert sheet["E1"].value == "2019-02-01"
        assert sheet.max_row == 97
    finally:
        workbook.close()


def test_save_pivot_workbook_respects_overwrite_flag(tmp_path: Path) -> None:
    result = _stream_result()
    filtered = result.df_stream.loc[
        (result.df_stream["region_code"].astype("object") == "0105a_region_001")
        & (result.df_stream["measure"].astype("object") == "debt")
        & (result.df_stream["currency_scope"].astype("object") == "total")
    ].copy()
    pivot_set = build_cbr_0105a_debt_corp_pivot_set(
        filtered,
        Cbr0105ADebtCorpPivotRequest(
            row_dimension="industry_code",
            column_dimension="report_date",
        ),
        stream_result=result,
    )
    store = LocalFileStore(root=str(tmp_path))
    request = Cbr0105ADebtCorpPivotWorkbookRequest(out_path="result.xlsx")
    save_cbr_0105a_debt_corp_pivot_workbook(pivot_set, request, filestore=store)
    with pytest.raises(FileExistsError):
        save_cbr_0105a_debt_corp_pivot_workbook(pivot_set, request, filestore=store)


def test_parser_reports_negative_values_as_validation_error() -> None:
    content = _workbook_bytes(
        report_date="01.06.2026",
        kemerovo_name="Кемеровская область - Кузбасс",
    )
    workbook = load_workbook(BytesIO(content))
    try:
        workbook[workbook.sheetnames[0]]["B4"] = -1
        buffer = BytesIO()
        workbook.save(buffer)
    finally:
        workbook.close()
    parsed = parse_cbr_0105a_debt_corp_excel_bytes(
        buffer.getvalue(),
        source_name="01_05_A_Debt_corp_20260601.xlsx",
    )
    errors = [issue for issue in parsed.validation_issues if issue.severity == "error"]
    assert "negative_values" in [issue.code for issue in errors]


def test_pivot_partial_stream_policy_is_explicit() -> None:
    result = _stream_result()
    partial = replace(
        result,
        failures=(
            Cbr0105ADebtCorpSourceFailure(
                stage="download",
                source_id="missing",
                url="https://example.invalid/missing.xlsx",
                source_name="missing.xlsx",
                report_date="2026-07-01",
                error="test",
            ),
        ),
    )
    filtered = partial.df_stream.loc[
        (partial.df_stream["region_code"].astype("object") == "0105a_region_001")
        & (partial.df_stream["measure"].astype("object") == "debt")
        & (partial.df_stream["currency_scope"].astype("object") == "total")
    ].copy()
    with pytest.raises(Cbr0105ADebtCorpPivotError) as exc_info:
        build_cbr_0105a_debt_corp_pivot_set(
            filtered,
            Cbr0105ADebtCorpPivotRequest(
                row_dimension="industry_code",
                column_dimension="report_date",
            ),
            stream_result=partial,
        )
    assert exc_info.value.code == "PIVOT_PARTIAL_STREAM"

    pivot_set = build_cbr_0105a_debt_corp_pivot_set(
        filtered,
        Cbr0105ADebtCorpPivotRequest(
            row_dimension="industry_code",
            column_dimension="report_date",
            require_complete_stream=False,
        ),
        stream_result=partial,
    )
    assert pivot_set.is_partial
