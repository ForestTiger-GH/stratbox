from pathlib import Path

import openpyxl

from stratbox.macrobanks.cbr_sors_restoration.contracts import (
    SorsPivotRequest,
    SorsWorkbookRequest,
)
from stratbox.macrobanks.cbr_sors_restoration.export import export_sors_workbook
from stratbox.macrobanks.cbr_sors_restoration.pivots import build_sors_pivot


def test_region_pivot_has_classes_as_rows(small_result) -> None:
    pivot = build_sors_pivot(small_result, SorsPivotRequest(region_name='Регион 1'))
    assert pivot.orientation == 'REGION_TO_CLASSES'
    assert tuple(pivot.table['class_code']) == ('01', '16')


def test_class_pivot_has_regions_as_rows(small_result) -> None:
    pivot = build_sors_pivot(small_result, SorsPivotRequest(class_code='16'))
    assert pivot.orientation == 'CLASS_TO_REGIONS'
    assert tuple(pivot.table['region_code']) == ('r1', 'r2')


def test_workbook_export_uses_pivot_contract(tmp_path: Path, small_result) -> None:
    output = tmp_path / 'sors.xlsx'
    exported = export_sors_workbook(
        small_result,
        SorsWorkbookRequest(
            output,
            pivots=(SorsPivotRequest(region_name='Регион 1'),),
        ),
    )
    assert exported.output_path == output
    assert exported.sha256
    workbook = openpyxl.load_workbook(output, read_only=True)
    assert 'Regional_OKVED2' in workbook.sheetnames
    assert any(name.startswith('Pivot_1_') for name in workbook.sheetnames)


def test_crosswalk_workbook_exports_primary_facts_bounds_and_mapping(
    tmp_path: Path,
    small_result,
) -> None:
    import pandas as pd

    from stratbox.macrobanks.cbr_sors_restoration.results import SorsCrosswalkResult

    primary = small_result.regional_okved2_grid.copy()
    primary['bounds_certified'] = True
    primary['value_identified'] = primary['value'].notna()
    primary['is_final_accepted'] = primary['value'].notna()
    primary['is_benchmark_estimate'] = False
    primary['benchmark_value'] = None
    primary['evidence_profile'] = 'core'
    primary['scenario_ids'] = [('core',)] * len(primary)
    facts = primary[primary['is_final_accepted'].astype(bool)].copy()
    mapping = pd.DataFrame(
        [{'scenario_id': 'core', 'atom_code': 'a', 'class_code': '01'}]
    )
    relations = pd.DataFrame(
        [{'scenario_id': 'core', 'relation_id': 'r1'}]
    )
    empty = pd.DataFrame()
    result = SorsCrosswalkResult(
        regional_okved2_grid=primary,
        crosswalk_bounds_grid=primary,
        crosswalk_facts_grid=facts,
        scenario_bounds_grid=primary.assign(scenario_id='core'),
        mapping_edges_grid=mapping,
        relations_grid=relations,
        constraints_grid=empty,
        variables_grid=empty,
        derivations_grid=empty,
        diagnostics_grid=empty,
        solver_runs_grid=empty,
        conflicts_grid=empty,
        audit_grid=pd.DataFrame([{'key': 'status', 'value': 'OPTIMAL'}]),
        crosswalk_run_id='test',
        status='OPTIMAL',
        _strict_result=small_result,
    )
    output = tmp_path / 'crosswalk.xlsx'
    exported = export_sors_workbook(
        result,
        SorsWorkbookRequest(
            output,
            include_crosswalk_bounds=True,
            include_crosswalk_mapping=True,
            pivots=(SorsPivotRequest(region_name='Регион 1'),),
        ),
    )

    workbook = openpyxl.load_workbook(output, read_only=True)
    assert exported.file_size > 0
    assert 'Crosswalk_Facts' in workbook.sheetnames
    assert 'Crosswalk_Bounds' in workbook.sheetnames
    assert 'Crosswalk_Scenarios' in workbook.sheetnames
    assert 'Crosswalk_Edges' in workbook.sheetnames
    assert 'Crosswalk_Relations' in workbook.sheetnames
    assert any(name.startswith('Pivot_1_') for name in workbook.sheetnames)
